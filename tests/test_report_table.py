from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ficha.report.naming import PLACEHOLDER_BASENAME, delivery_basename, surname_of
from ficha.report.table import FICHA_COLUMNS, fichas_to_frame, write_table
from ficha.schema import Confianca, Evidencia, Ficha


def _ficha(arquivo: str, limitacao: str | None, conf: Confianca | None) -> Ficha:
    return Ficha(
        arquivo=arquivo,
        problema="Prever readmissão hospitalar em 30 dias.",
        dados="MIMIC-IV, 2008–2019, 431.231 internações.",
        metodo="LightGBM com TF-IDF de notas de alta.",
        metrica="AUROC e AUPRC.",
        limitacao=limitacao,
        evidencia=Evidencia(
            trecho="We use the MIMIC-IV database, covering 2008 to 2019.", pagina=2
        ),
        confianca=conf,
    )


@pytest.fixture
def fichas() -> list[Ficha]:
    return [
        _ficha("artigo_01.pdf", "Dados de uma única instituição.", Confianca.ALTA),
        _ficha("artigo_02.pdf", None, Confianca.MEDIA),
        _ficha("artigo_03.pdf", None, None),
    ]


def test_frame_columns_follow_section_3(fichas: list[Ficha]) -> None:
    frame = fichas_to_frame(fichas)
    assert tuple(frame.columns) == FICHA_COLUMNS
    assert list(frame.columns) == list(fichas[0].to_row())
    assert frame["evidencia_pagina"].dtype == "int64"
    assert len(frame) == 3


def test_frame_extras_come_after_ficha_columns(fichas: list[Ficha]) -> None:
    extras = {"artigo_02.pdf": {"motivos_confianca": ["trecho parcial", "instável"]}}
    frame = fichas_to_frame(fichas, extras)
    assert list(frame.columns) == [*FICHA_COLUMNS, "motivos_confianca"]
    assert frame.loc[1, "motivos_confianca"] == "trecho parcial; instável"
    assert frame.loc[0, "motivos_confianca"] is None


def test_write_table_csv_and_xlsx(tmp_path: Path, fichas: list[Ficha]) -> None:
    out = write_table(fichas, tmp_path / "out", basename="AtividadeI_Teste")
    assert out.csv.name == "AtividadeI_Teste.csv"
    assert out.xlsx.name == "AtividadeI_Teste.xlsx"
    assert out.csv.exists() and out.xlsx.exists()

    raw = out.csv.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf"), "CSV precisa de BOM para o Excel pt-BR"
    with out.csv.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    assert list(rows[0]) == list(FICHA_COLUMNS)
    assert rows[1]["limitacao"] == ""  # None vira célula vazia, não "None"
    assert "None" not in out.csv.read_text(encoding="utf-8-sig")
    assert rows[2]["confianca"] == ""
    assert rows[1]["dados"] == "MIMIC-IV, 2008–2019, 431.231 internações."

    wb = load_workbook(out.xlsx)
    ws = wb.active
    assert ws is not None
    header = [c.value for c in ws[1]]
    assert header == list(FICHA_COLUMNS)
    assert ws.cell(row=1, column=1).font.bold
    assert ws.freeze_panes == "B2"
    limit_col = header.index("limitacao") + 1
    assert ws.cell(row=3, column=limit_col).value is None
    assert ws.cell(row=2, column=limit_col).alignment.wrap_text


def test_delivery_basename() -> None:
    assert delivery_basename(["Ana Souza", "João da Conceição"]) == "AtividadeI_Souza_Conceicao"
    assert delivery_basename([]) == PLACEHOLDER_BASENAME
    assert delivery_basename(["<<INTEGRANTES: Nome Sobrenome>>"]) == PLACEHOLDER_BASENAME
    assert surname_of("  Maria  Araújo ") == "Araujo"
    with pytest.raises(ValueError):
        surname_of("   ")

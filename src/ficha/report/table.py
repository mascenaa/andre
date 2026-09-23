"""Tabela comparativa dos artigos (Seção 5, entregável 2): CSV e XLSX.

Uma linha por artigo, com **todos** os campos da ficha na ordem da Seção 3
(:meth:`ficha.schema.Ficha.to_row`). Colunas extras de auditoria (por exemplo, os motivos
da confiança atribuída) são opcionais e vêm sempre **depois** das colunas da ficha, para que
quem abre a planilha veja primeiro exatamente o que o enunciado pede.

Decisões de formato:

- CSV em UTF-8 **com BOM** (``utf-8-sig``): sem o BOM, o Excel em pt-BR abre "ção" como
  "Ã§Ã£o". O separador continua sendo vírgula (padrão do formato).
- ``limitacao = null`` vira **célula vazia**, nunca a string ``"None"`` — na tabela, vazio é a
  forma legível de "os autores não declararam".
- XLSX com cabeçalho em negrito, quebra de linha, larguras fixas razoáveis e painel
  congelado (cabeçalho + coluna ``arquivo``), para leitura direta pelo coordenador.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ficha.report.naming import PLACEHOLDER_BASENAME
from ficha.schema import Ficha

FICHA_COLUMNS: tuple[str, ...] = (
    "arquivo",
    "problema",
    "dados",
    "metodo",
    "metrica",
    "limitacao",
    "evidencia_trecho",
    "evidencia_pagina",
    "confianca",
)
"""Colunas da ficha, na ordem da Seção 3 (idêntica a :meth:`Ficha.to_row`)."""

_COLUMN_WIDTHS: dict[str, int] = {
    "arquivo": 22,
    "problema": 45,
    "dados": 45,
    "metodo": 45,
    "metrica": 35,
    "limitacao": 45,
    "evidencia_trecho": 60,
    "evidencia_pagina": 10,
    "confianca": 11,
}
_DEFAULT_EXTRA_WIDTH = 50


@dataclass(frozen=True, slots=True)
class TableOutputs:
    """Caminhos dos arquivos gerados por :func:`write_table`."""

    csv: Path
    xlsx: Path


def _cell(value: Any) -> Any:
    """Normaliza um valor de coluna extra: listas viram texto separado por ``; ``."""
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, Sequence):
        return "; ".join(str(v) for v in value)
    return value


def fichas_to_frame(
    fichas: Sequence[Ficha],
    extras: Mapping[str, Mapping[str, Any]] | None = None,
) -> pd.DataFrame:
    """Converte fichas em um ``DataFrame`` com as colunas da Seção 3 na ordem exata.

    Args:
        fichas: fichas finais (normalmente já com ``confianca`` atribuída pela auditoria).
        extras: colunas de auditoria opcionais, por arquivo — ex.:
            ``{"artigo_01.pdf": {"motivos_confianca": ["trecho não encontrado"]}}``.
            Listas viram texto separado por ``"; "``. Artigos sem extras ficam com célula vazia.

    Returns:
        Um ``DataFrame`` com uma linha por ficha, na ordem recebida.

    """
    rows = [f.to_row() for f in fichas]
    extra_cols: list[str] = []
    if extras:
        for per_file in extras.values():
            for col in per_file:
                if col not in extra_cols and col not in FICHA_COLUMNS:
                    extra_cols.append(col)
        for row in rows:
            per_file = extras.get(str(row["arquivo"]), {})
            for col in extra_cols:
                row[col] = _cell(per_file.get(col))
    frame = pd.DataFrame(rows, columns=[*FICHA_COLUMNS, *extra_cols])
    # Ausência é sempre ``None`` (nunca ``NaN``), para CSV/XLSX gravarem célula vazia.
    text_cols = [c for c in frame.columns if c != "evidencia_pagina"]
    frame[text_cols] = frame[text_cols].astype(object).where(frame[text_cols].notna(), None)
    # Página é inteira; sem isso, uma coluna vazia viraria float.
    if not frame.empty:
        frame["evidencia_pagina"] = frame["evidencia_pagina"].astype("int64")
    return frame


def _style_xlsx(path: Path, columns: Sequence[str]) -> None:
    """Aplica cabeçalho em negrito, quebra de linha, larguras e painel congelado."""
    wb = load_workbook(path)
    ws = wb.active
    assert ws is not None  # planilha recém-criada pelo pandas sempre tem aba ativa
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E7E6E6")
    wrap = Alignment(wrap_text=True, vertical="top")
    for idx, col in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = _COLUMN_WIDTHS.get(col, _DEFAULT_EXTRA_WIDTH)
        header = ws.cell(row=1, column=idx)
        header.font = header_font
        header.fill = header_fill
        header.alignment = wrap
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_table(
    fichas: Sequence[Ficha],
    out_dir: Path,
    basename: str = PLACEHOLDER_BASENAME,
    extras: Mapping[str, Mapping[str, Any]] | None = None,
) -> TableOutputs:
    """Grava a tabela comparativa em CSV (UTF-8 com BOM) e XLSX formatado.

    Args:
        fichas: fichas finais, uma por artigo.
        out_dir: diretório de saída (criado se não existir).
        basename: nome-base dos arquivos, ``AtividadeI_<sobrenomes>`` (ver
            :func:`ficha.report.naming.delivery_basename`).
        extras: colunas de auditoria opcionais (ver :func:`fichas_to_frame`).

    Returns:
        Os caminhos do CSV e do XLSX gerados.

    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    frame = fichas_to_frame(fichas, extras)
    csv_path = out_dir / f"{basename}.csv"
    xlsx_path = out_dir / f"{basename}.xlsx"
    frame.to_csv(csv_path, index=False, encoding="utf-8-sig")
    frame.to_excel(xlsx_path, index=False, sheet_name="fichas", engine="openpyxl")
    _style_xlsx(xlsx_path, list(frame.columns))
    return TableOutputs(csv=csv_path, xlsx=xlsx_path)
